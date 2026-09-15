"""Paso 6 — Ensamble del .pptx. Estilo visual calcado del deck real de Rebold
("Plaza_Mayor_Propuesta_Rebold.pptx"): fondo casi negro, tarjetas oscuras con
esquinas redondeadas, kicker rosa en mayúsculas, titulares en Arial Black,
cifras grandes, y numeración de página en la esquina.

Paso 7 (obligatorio, sin excepción): todo archivo que sale de aquí lleva una
marca de BORRADOR — PENDIENTE DE REVISIÓN visible en cada slide. Esta función
nunca envía nada a nadie, solo genera el archivo para que el ejecutivo lo
revise y ajuste manualmente, especialmente la slide de inversión.
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --- Paleta real de Rebold (extraída del deck de referencia) ---
BG = RGBColor(0x0B, 0x0B, 0x0C)
CARD = RGBColor(0x13, 0x13, 0x16)
CARD_ALT = RGBColor(0x17, 0x15, 0x1A)
ICON_BG = RGBColor(0x1D, 0x1D, 0x21)
CHERRY = RGBColor(0xCD, 0x2B, 0x53)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x9A, 0x9A, 0xA2)
MUTED_DARK = RGBColor(0x5A, 0x5A, 0x62)

FONT_BODY = "Arial"
FONT_HEAD = "Arial Black"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = SLIDE_W - 2 * MARGIN


def _fondo(slide, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _texto(
    slide, left, top, width, height, texto: str, *,
    size: int = 14, color: RGBColor = WHITE, bold: bool = False,
    align=PP_ALIGN.LEFT, font: str = FONT_BODY, italic: bool = False,
    anchor=MSO_ANCHOR.TOP, line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = texto
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return box


def _rect(slide, left, top, width, height, color: RGBColor, *, radius: float = 0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        shape.adjustments[0] = radius
    except IndexError:
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _circle(slide, left, top, diameter, color: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, diameter, diameter)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _agregar_notas(slide, texto: str) -> None:
    if texto:
        slide.notes_slide.notes_text_frame.text = texto


def _header(slide, kicker: str, titulo: str, subtitulo: str = "", *, titulo_size: int = 27) -> None:
    _texto(slide, MARGIN, Inches(0.38), CONTENT_W, Inches(0.3), kicker.upper(),
           size=11.5, color=CHERRY, bold=True)
    alto_titulo = Inches(0.85) if len(titulo) < 55 else Inches(1.15)
    _texto(slide, MARGIN, Inches(0.72), CONTENT_W, alto_titulo, titulo,
           size=titulo_size, color=WHITE, bold=True, font=FONT_HEAD)
    if subtitulo:
        top = Inches(0.72) + alto_titulo + Inches(0.05)
        _texto(slide, MARGIN, top, CONTENT_W, Inches(0.7), subtitulo,
               size=13, color=MUTED)


def _footer(slide, pagina: int) -> None:
    _texto(slide, MARGIN, Inches(7.08), Inches(6), Inches(0.3),
           "BORRADOR — PENDIENTE DE REVISIÓN INTERNA", size=9, color=MUTED_DARK, bold=True)
    _texto(slide, SLIDE_W - Inches(1.0), Inches(7.08), Inches(0.6), Inches(0.3),
           f"{pagina:02d}", size=10, color=MUTED_DARK, align=PP_ALIGN.RIGHT)


def _card_grid(slide, cards: list[dict[str, str]], *, top=Inches(1.9), bottom=Inches(6.85)) -> None:
    """cards: [{"badge": "01", "titulo": "...", "cuerpo": "...", "tag": "..." (opcional)}]"""
    n = len(cards)
    gap = Inches(0.28)
    w = (CONTENT_W - gap * (n - 1)) / n if n else CONTENT_W
    h = bottom - top
    for i, card in enumerate(cards):
        x = MARGIN + i * (w + gap)
        bg = CARD if i % 2 == 0 else CARD_ALT
        _rect(slide, x, top, w, h, bg)
        icon_bg = CHERRY if i % 2 == 1 else ICON_BG
        d = Inches(0.5)
        _circle(slide, x + Inches(0.3), top + Inches(0.3), d, icon_bg)
        _texto(slide, x + Inches(0.3), top + Inches(0.3), d, d, card.get("badge", ""),
               size=14, color=WHITE if icon_bg == CHERRY else MUTED, bold=True,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=FONT_HEAD)
        _texto(slide, x + Inches(0.3), top + Inches(1.0), w - Inches(0.6), Inches(0.6),
               card["titulo"], size=15.5, color=WHITE, bold=True, font=FONT_HEAD)
        _texto(slide, x + Inches(0.3), top + Inches(1.55), w - Inches(0.6), h - Inches(2.0),
               card.get("cuerpo", ""), size=11, color=MUTED)
        if card.get("tag"):
            tag_h = Inches(0.4)
            tag_y = top + h - tag_h - Inches(0.25)
            _rect(slide, x + Inches(0.3), tag_y, w - Inches(0.6), tag_h, ICON_BG, radius=0.3)
            _texto(slide, x + Inches(0.3), tag_y, w - Inches(0.6), tag_h, card["tag"],
                   size=10, color=CHERRY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _stat_list(slide, x, y, w, items: list[str]) -> None:
    row_h = Inches(0.42)
    for i, item in enumerate(items):
        yy = y + i * row_h
        _texto(slide, x, yy, Inches(0.3), row_h, "✱", size=13, color=CHERRY,
               anchor=MSO_ANCHOR.MIDDLE)
        _texto(slide, x + Inches(0.32), yy, w - Inches(0.32), row_h, item,
               size=13, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, font=FONT_HEAD)


def _investment_card(slide, x, y, w, h, *, kicker: str, titulo: str, precio: str,
                      subtexto: str = "", accent: bool = False) -> None:
    bg = CARD_ALT if accent else CARD
    _rect(slide, x, y, w, h, bg)
    icon_bg = CHERRY if accent else ICON_BG
    d = Inches(0.55)
    _circle(slide, x + Inches(0.32), y + Inches(0.32), d, icon_bg)
    _texto(slide, x + Inches(1.05), y + Inches(0.32), w - Inches(1.35), d, kicker.upper(),
           size=10.5, color=CHERRY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _texto(slide, x + Inches(0.32), y + Inches(1.05), w - Inches(0.64), Inches(0.55),
           titulo, size=16, color=WHITE, bold=True, font=FONT_HEAD)

    # Si el texto trae una explicación entre paréntesis (ej. la regla escalonada
    # de Activation Hub), se separa: el precio queda grande y corto, la
    # explicación baja como texto secundario para no saturar la tarjeta.
    detalle_regla = ""
    if " (" in precio and precio.endswith(")"):
        precio, _, resto = precio.partition(" (")
        detalle_regla = resto[:-1]

    color_precio = WHITE if ("pendiente" not in precio.lower()) else MUTED
    size_precio = 24 if len(precio) < 34 else 17
    _texto(slide, x + Inches(0.32), y + Inches(1.65), w - Inches(0.64), Inches(0.85),
           precio, size=size_precio, color=color_precio, bold=True, font=FONT_HEAD)

    y_subtexto = y + Inches(2.45)
    if detalle_regla:
        _texto(slide, x + Inches(0.32), y_subtexto, w - Inches(0.64), Inches(0.6),
               detalle_regla, size=10.5, color=MUTED)
        y_subtexto += Inches(0.55)
    if subtexto:
        _texto(slide, x + Inches(0.32), y_subtexto, w - Inches(0.64), h - (y_subtexto - y),
               subtexto, size=10.5, color=MUTED)


@dataclass
class ContenidoPropuesta:
    cliente: str
    fecha: str
    tipo_relacion: str
    narrativa_contexto: str
    narrativa_insight: str
    dolor_principal_cita: str
    recomendaciones: list[dict[str, Any]]  # de diagnosis_a_dict()["recomendaciones"]
    camino_de_cuenta: list[str]
    oportunidades_futuras: list[str]
    lineas_inversion: list[dict[str, str]]  # [{"servicio","modalidad","estado","texto_cliente","nota_interna"}]
    fuente_pricing: str
    contacto_nombre: str = ""
    contacto_email: str = ""


def construir_pptx(contenido: ContenidoPropuesta, ruta_salida: str) -> str:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    layout_vacio = prs.slide_layouts[6]
    pagina = 0

    def nueva_slide(color=BG):
        nonlocal pagina
        pagina += 1
        s = prs.slides.add_slide(layout_vacio)
        _fondo(s, color)
        return s

    tipo_label = "Cliente actual" if contenido.tipo_relacion == "cliente_actual" else "Prospecto nuevo"

    # --- 1. Portada ---
    s = nueva_slide()
    _texto(s, MARGIN, Inches(1.15), Inches(6), Inches(0.3), "REBOLD BY ISPD", size=12, color=MUTED_DARK, bold=True)
    _texto(s, MARGIN, Inches(1.55), Inches(8), Inches(0.45), "PROPUESTA COMERCIAL", size=15, color=CHERRY, bold=True)
    _texto(s, MARGIN, Inches(1.95), CONTENT_W, Inches(1.5), contenido.cliente, size=48, color=WHITE, bold=True, font=FONT_HEAD)
    tagline = contenido.narrativa_insight or "Diagnóstico, narrativa y camino de cuenta construidos a partir de la conversación real con el cliente."
    _texto(s, MARGIN, Inches(3.55), Inches(9.5), Inches(0.7), tagline, size=14, color=MUTED)

    tiles = [
        (str(len(contenido.recomendaciones)), "SERVICIO(S) RECOMENDADO(S)"),
        (tipo_label.upper(), "TIPO DE RELACIÓN"),
        (contenido.fecha, "FECHA DE LA REUNIÓN"),
    ]
    tile_w = Inches(2.6)
    for i, (big, small) in enumerate(tiles):
        x = MARGIN + i * (tile_w + Inches(0.25))
        _rect(s, x, Inches(4.7), tile_w, Inches(1.25), CARD)
        size_big = 22 if len(big) <= 6 else (16 if len(big) <= 12 else 13)
        _texto(s, x + Inches(0.25), Inches(4.82), tile_w - Inches(0.5), Inches(0.65), big,
               size=size_big, color=CHERRY, bold=True, font=FONT_HEAD)
        _texto(s, x + Inches(0.25), Inches(5.55), tile_w - Inches(0.5), Inches(0.35), small,
               size=9.5, color=MUTED, bold=True)
    _texto(s, MARGIN, Inches(6.35), Inches(6), Inches(0.3), "A Post Digital Marketing & Communication Company",
           size=9, color=MUTED_DARK)
    _footer(s, pagina)
    _agregar_notas(s, "Borrador generado automáticamente a partir de la nota de reunión. Revisar por completo antes de compartir con el cliente.")

    # --- 2. Contexto y dolor validado ---
    s = nueva_slide()
    _header(s, "Contexto y oportunidad", "Esto fue lo que dijo el cliente.", contenido.narrativa_contexto)
    if contenido.dolor_principal_cita:
        _rect(s, MARGIN, Inches(3.1), CONTENT_W, Inches(2.9), CARD)
        _rect(s, MARGIN, Inches(3.1), Inches(0.08), Inches(2.9), CHERRY, radius=0.0)
        _texto(s, MARGIN + Inches(0.5), Inches(3.5), CONTENT_W - Inches(1.0), Inches(1.8),
               f"“{contenido.dolor_principal_cita}”", size=24, color=WHITE, bold=True,
               italic=True, font=FONT_BODY, line_spacing=1.15)
        primer_servicio = contenido.recomendaciones[0]["nombre"] if contenido.recomendaciones else "Diagnóstico"
        tag_y = Inches(5.55)
        _rect(s, MARGIN + Inches(0.5), tag_y, Inches(4.5), Inches(0.42), ICON_BG, radius=0.3)
        _texto(s, MARGIN + Inches(0.5), tag_y, Inches(4.5), Inches(0.42), f"Dolor validado · {primer_servicio}",
               size=10.5, color=CHERRY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _agregar_notas(s, "Verifica que esta cita aparezca literalmente en la nota original antes de enviar.")
    else:
        _rect(s, MARGIN, Inches(3.1), CONTENT_W, Inches(2.0), CARD)
        _texto(s, MARGIN + Inches(0.4), Inches(3.4), CONTENT_W - Inches(0.8), Inches(1.4),
               "Primera conversación exploratoria. Se abre con un diagnóstico objetivo (Rebold Audit) "
               "en vez de asumir un dolor que el cliente todavía no ha dicho.",
               size=16, color=MUTED)
        _agregar_notas(s, "Sin cita textual: esta propuesta abre con Rebold Audit por defecto (prospecto nuevo sin dolor explícito).")
    _footer(s, pagina)

    # --- 3. Diagnóstico / servicios recomendados ---
    s = nueva_slide()
    _header(s, "Diagnóstico Rebold", "Lo que recomendamos y por qué.",
            "Máximo dos servicios por propuesta inicial — el resto queda como oportunidad futura interna.")
    cards = []
    for i, rec in enumerate(contenido.recomendaciones):
        cards.append({
            "badge": chr(65 + i),
            "titulo": rec["nombre"],
            "cuerpo": rec["pitch"],
            "tag": f"Dolor: “{rec['dolor_citado'][:70]}{'…' if len(rec['dolor_citado']) > 70 else ''}”" if rec["dolor_citado"] else "Apertura por defecto",
        })
    _card_grid(s, cards, top=Inches(2.15))
    _footer(s, pagina)
    notas_diag = "Camino de cuenta / cross-sell (NO mostrar en la slide visible):\n" + "\n".join(contenido.camino_de_cuenta)
    if contenido.oportunidades_futuras:
        notas_diag += "\n\nOportunidades futuras (no incluidas en esta propuesta): " + ", ".join(contenido.oportunidades_futuras)
    _agregar_notas(s, notas_diag)

    # --- 4. Proof case (una slide por servicio recomendado) ---
    for rec in contenido.recomendaciones:
        caso = rec.get("proof_case") or {}
        s = nueva_slide()
        if caso.get("cliente"):
            _header(s, f"Caso de referencia · {rec['nombre']}",
                    f"{caso['cliente']} · {caso.get('industria', '')}", caso.get("resumen", ""))
            cifras = caso.get("cifras", [])
            if cifras:
                _rect(s, MARGIN, Inches(2.6), CONTENT_W, Inches(3.6), CARD)
                _stat_list(s, MARGIN + Inches(0.4), Inches(2.95), CONTENT_W - Inches(0.8), cifras)
        else:
            _header(s, f"Caso de referencia · {rec['nombre']}", "Sin caso de referencia disponible todavía.")
            _rect(s, MARGIN, Inches(2.6), CONTENT_W, Inches(1.6), CARD)
            _texto(s, MARGIN + Inches(0.4), Inches(2.9), CONTENT_W - Inches(0.8), Inches(1.0),
                   "No hay un caso de prueba registrado para este servicio en la skill de Rebold. "
                   "No se fabricó un caso — actualizar references/casos-prueba.md antes del próximo pitch.",
                   size=13, color=MUTED)
        _footer(s, pagina)

    # --- 5. Inversión (una tarjeta por servicio recomendado) ---
    s = nueva_slide()
    _header(s, "Inversión", "Lo que cuesta cada servicio recomendado.")
    n = len(contenido.lineas_inversion)
    gap = Inches(0.3)
    card_w = (CONTENT_W - gap * (n - 1)) / n if n else CONTENT_W
    card_h = Inches(3.6)
    top = Inches(2.2)
    for i, linea in enumerate(contenido.lineas_inversion):
        x = MARGIN + i * (card_w + gap)
        _investment_card(
            s, x, top, card_w, card_h,
            kicker=linea["modalidad"], titulo=linea["servicio"], precio=linea["texto_cliente"],
            subtexto="Cifra definida en la tabla de pricing de liderazgo." if linea["estado"] != "pendiente"
                     else "Debe cerrarse con liderazgo Rebold antes de enviar al cliente.",
            accent=(i % 2 == 1),
        )
    _texto(s, MARGIN, top + card_h + Inches(0.25), CONTENT_W, Inches(0.5),
           f"Fuente: {contenido.fuente_pricing}. Cifras 'pendiente' se definen exclusivamente con liderazgo Rebold — no editar sin su aprobación.",
           size=10, color=MUTED_DARK)
    _footer(s, pagina)
    notas_inversion = "\n".join(l.get("nota_interna", "") for l in contenido.lineas_inversion if l.get("nota_interna"))
    _agregar_notas(s, "Notas internas de pricing (NO leer en voz alta frente al cliente):\n" + notas_inversion)

    # --- 6. Próximos pasos ---
    s = nueva_slide()
    _header(s, "Próximos pasos", "De la propuesta al primer entregable.")
    pasos = [
        {"badge": "01", "titulo": "Aprobación", "cuerpo": "El cliente aprueba el alcance y confirma fecha de inicio."},
        {"badge": "02", "titulo": "Kick-off", "cuerpo": "Rebold alinea objetivos, accesos y responsables por frente."},
        {"badge": "03", "titulo": "Ejecución", "cuerpo": "Se ejecuta el/los servicio(s) recomendado(s) según el modelo de trabajo elegido."},
        {"badge": "04", "titulo": "Resultados", "cuerpo": "Entrega, medición y siguiente paso de cuenta (cross-sell)."},
    ]
    _card_grid(s, pasos, top=Inches(2.15))
    _footer(s, pagina)

    # --- 7. Cierre ---
    s = nueva_slide()
    _texto(s, MARGIN, Inches(2.9), CONTENT_W, Inches(0.8), "REBOLD", size=40, color=CHERRY, bold=True, font=FONT_HEAD, align=PP_ALIGN.CENTER)
    _texto(s, MARGIN, Inches(3.75), CONTENT_W, Inches(0.5), "GROWTH FOR TODAY AND TOMORROW",
           size=13, color=MUTED, align=PP_ALIGN.CENTER)
    contacto = " · ".join(x for x in [contenido.contacto_nombre, contenido.contacto_email] if x)
    if contacto:
        _texto(s, MARGIN, Inches(4.4), CONTENT_W, Inches(0.4), f"Preparado para {contacto}",
               size=11, color=MUTED_DARK, align=PP_ALIGN.CENTER)
    _footer(s, pagina)

    prs.save(ruta_salida)
    return ruta_salida
